"""
Convierte data/claims_etiquetado.csv en data/claims_etiquetado.xlsx con:
  - Leeme: origen del dataset y cómo regenerarlo
  - claims: los 240 registros (filtro y encabezado fijo)
  - diccionario_datos: descripción de cada campo
  - guia_etiquetado: reglas de anotación para casos frontera
  - resumen: conteos por clase, partición y dificultad (fórmulas COUNTIFS)

Uso:  python src/exportar_excel.py
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parents[1]
CSV = RAIZ / "data" / "claims_etiquetado.csv"
XLSX = RAIZ / "data" / "claims_etiquetado.xlsx"

AZUL = "1F4E79"
F = "Arial"
HDR_FILL = PatternFill("solid", fgColor=AZUL)
HDR_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
TXT = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CLASES = ["NUTRIMENTAL", "SALUD", "AUSENCIA", "MARKETING"]

DICCIONARIO = [
    ("id_claim", "Texto", "Identificador único del claim (CLM-0001 ... CLM-0240).", "CLM-0007"),
    ("texto_claim", "Texto", "Claim tal como aparecería en el empaque. Es la entrada del modelo.", "Sin azúcares añadidos"),
    ("categoria", "Categórico (etiqueta objetivo)", "Tipo de claim según la guía de etiquetado: NUTRIMENTAL, SALUD, AUSENCIA o MARKETING.", "NUTRIMENTAL"),
    ("subtipo", "Categórico", "Subtipo fino usado para el análisis de errores (p. ej., comparativa, reduccion_riesgo, aditivo).", "contenido_bajo_o_sin"),
    ("referencia_normativa", "Texto", "Marco regulatorio que inspira la redacción del claim. Sirve de contexto; no es un dictamen de cumplimiento.", "Reg. (CE) 1924/2006, Anexo; ..."),
    ("dificultad", "Categórico", "estandar = caso típico; ambiguo = caso frontera entre dos clases según la guía.", "ambiguo"),
    ("origen_registro", "Categórico", "redaccion_manual = escrito a mano; plantilla = generado por combinación de plantillas con semilla fija.", "redaccion_manual"),
    ("n_palabras", "Entero", "Número de palabras del claim (separadas por espacio).", "3"),
    ("particion", "Categórico", "dev (30 %, estratificado) para elegir plantilla de hipótesis y umbral; test (70 %) para las métricas finales.", "test"),
]

GUIA = [
    ("SALUD", "El claim relaciona un alimento, nutrimento o componente con una función del cuerpo, la salud o la reducción de riesgo de enfermedad. Tiene prioridad aunque también mencione el contenido.",
     "Fuente de calcio, que ayuda a mantener huesos fuertes → SALUD"),
    ("NUTRIMENTAL", "Declara presencia, cantidad, ausencia o comparación de un nutrimento (energía, grasa, grasas saturadas o trans, azúcares, sodio/sal, fibra, proteína, vitaminas, minerales, colesterol).",
     "Sin azúcares añadidos → NUTRIMENTAL (azúcares es un nutrimento)"),
    ("AUSENCIA", "Declara la ausencia de algo que no es un nutrimento: alérgeno o sustancia de intolerancia, aditivo, ingrediente específico o práctica de producción.",
     "Sin edulcorantes artificiales → AUSENCIA (aditivo, no nutrimento)"),
    ("MARKETING", "Descriptor de calidad, origen, proceso, sabor, textura, conveniencia, sostenibilidad o 'natural', sin referencia a nutrimentos ni a la salud.",
     "Ligero y crujiente → MARKETING ('ligero' describe textura, no energía)"),
]


def encabezado(ws, fila, valores, anchos=None):
    for j, v in enumerate(valores, start=1):
        c = ws.cell(row=fila, column=j, value=v)
        c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
    if anchos:
        for j, w in enumerate(anchos, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w


def main():
    df = pd.read_csv(CSV, encoding="utf-8")
    wb = Workbook()

    # Leeme -----------------------------------------------------------------
    ws = wb.active
    ws.title = "Leeme"
    ws["A1"] = "Dataset de claims de etiquetado de alimentos (ClaimCheck IA · Actividad 4)"
    ws["A1"].font = Font(name=F, size=13, bold=True, color=AZUL)
    filas = [
        ("Propósito", "Comparar modelos preentrenados de Hugging Face que clasifican claims de etiquetas de alimentos por tipo, para enrutar cada claim al checklist de revisión correcto."),
        ("Origen", "Dataset propio. Los claims NO provienen de etiquetas reales de clientes: se redactaron a partir de la estructura y redacción típica de los marcos UE (Reg. 1924/2006, 432/2012, 828/2014, 1169/2011), FDA (21 CFR 101) y NOM-051-SCFI/SSA1-2010."),
        ("Composición", f"{len(df)} claims · {(df.origen_registro == 'redaccion_manual').sum()} redactados a mano y {(df.origen_registro == 'plantilla').sum()} generados por plantillas con semilla 42 · {(df.dificultad == 'ambiguo').sum()} casos frontera marcados como 'ambiguo'."),
        ("Reproducibilidad", "python src/generar_dataset.py regenera el CSV idéntico (misma semilla, IDs y partición). python src/exportar_excel.py regenera este archivo."),
        ("Hojas", "claims · diccionario_datos · guia_etiquetado · resumen"),
        ("Límite", "Al ser redactado por el autor, puede ser más 'limpio' que las etiquetas reales: las métricas son un techo optimista y deben validarse con claims reales anonimizados."),
    ]
    for i, (k, v) in enumerate(filas, start=3):
        ws.cell(row=i, column=1, value=k).font = BOLD
        c = ws.cell(row=i, column=2, value=v)
        c.font, c.alignment = TXT, WRAP
        ws.row_dimensions[i].height = 44
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 105

    # claims ----------------------------------------------------------------
    ws = wb.create_sheet("claims")
    cols = list(df.columns)
    encabezado(ws, 1, cols, [11, 52, 14, 24, 48, 11, 17, 11, 10])
    for i, fila in enumerate(df.itertuples(index=False), start=2):
        for j, v in enumerate(fila, start=1):
            c = ws.cell(row=i, column=j, value=v.item() if hasattr(v, "item") else v)
            c.font, c.border = TXT, BOX
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(df) + 1}"

    # diccionario -----------------------------------------------------------
    ws = wb.create_sheet("diccionario_datos")
    encabezado(ws, 1, ["Campo", "Tipo", "Descripción", "Ejemplo"], [22, 26, 80, 30])
    for i, fila in enumerate(DICCIONARIO, start=2):
        for j, v in enumerate(fila, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.alignment, c.border = (BOLD if j == 1 else TXT), WRAP, BOX

    # guía ------------------------------------------------------------------
    ws = wb.create_sheet("guia_etiquetado")
    ws["A1"] = "Reglas de anotación (se aplican en este orden de prioridad)"
    ws["A1"].font = Font(name=F, size=12, bold=True, color=AZUL)
    encabezado(ws, 3, ["Clase", "Regla", "Caso frontera resuelto"], [16, 85, 55])
    for i, fila in enumerate(GUIA, start=4):
        for j, v in enumerate(fila, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.alignment, c.border = (BOLD if j == 1 else TXT), WRAP, BOX
        ws.row_dimensions[i].height = 48

    # resumen (fórmulas) ----------------------------------------------------
    ws = wb.create_sheet("resumen")
    n = len(df) + 1
    rng = lambda col: f"claims!${col}$2:${col}${n}"  # noqa: E731
    encabezado(ws, 1, ["Clase", "Total", "% del total", "dev", "test", "ambiguo", "redaccion_manual", "plantilla"],
               [16, 10, 12, 10, 10, 11, 18, 12])
    for i, clase in enumerate(CLASES, start=2):
        ws.cell(row=i, column=1, value=clase).font = BOLD
        ws.cell(row=i, column=2, value=f'=COUNTIFS({rng("C")},$A{i})')
        ws.cell(row=i, column=3, value=f"=IF($B$6=0,0,B{i}/$B$6)").number_format = "0.0%"
        ws.cell(row=i, column=4, value=f'=COUNTIFS({rng("C")},$A{i},{rng("I")},"dev")')
        ws.cell(row=i, column=5, value=f'=COUNTIFS({rng("C")},$A{i},{rng("I")},"test")')
        ws.cell(row=i, column=6, value=f'=COUNTIFS({rng("C")},$A{i},{rng("F")},"ambiguo")')
        ws.cell(row=i, column=7, value=f'=COUNTIFS({rng("C")},$A{i},{rng("G")},"redaccion_manual")')
        ws.cell(row=i, column=8, value=f'=COUNTIFS({rng("C")},$A{i},{rng("G")},"plantilla")')
    ws.cell(row=6, column=1, value="Total").font = BOLD
    for j in range(2, 9):
        col = get_column_letter(j)
        c = ws.cell(row=6, column=j, value=f"=SUM({col}2:{col}5)")
        c.font = BOLD
        if j == 3:
            c.number_format = "0.0%"
    for fila in ws.iter_rows(min_row=2, max_row=6, max_col=8):
        for c in fila:
            c.border = BOX
            if c.font != BOLD and c.column != 1:
                c.font = TXT
    ws["A8"] = "Control: el total debe coincidir con el número de filas de la hoja claims."
    ws["A8"].font = Font(name=F, size=9, italic=True)
    ws["A9"] = "Filas en claims"
    ws["B9"] = f'=COUNTA({rng("A")})'
    ws["A10"] = "¿Coincide?"
    ws["B10"] = '=IF(B9=B6,"OK","REVISAR")'
    for ref in ("A9", "A10"):
        ws[ref].font = BOLD
    for ref in ("B9", "B10"):
        ws[ref].font = TXT

    wb.save(XLSX)
    print(f"Escrito {XLSX}")


if __name__ == "__main__":
    main()
